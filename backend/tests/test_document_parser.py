from pathlib import Path

import pytest
from app.services.document_parser import (
    ParsedDocument,
    ParsedSection,
    _parse_docx_structure,
    _parse_markdown_structure,
    assign_parent_orders,
    match_heading_style,
    normalize_text,
    parse_document_structure,
    parse_document_text,
)
from pydantic import ValidationError


def test_normalize_text_trims_lines_and_collapses_blank_lines() -> None:
    assert normalize_text(" first  \r\n\r\n\r\n second ") == "first\n\nsecond"


def test_parse_document_text_reads_utf8_markdown(tmp_path: Path) -> None:
    document = tmp_path / "sample.md"
    document.write_text("# 标题\n\n内容", encoding="utf-8")

    assert parse_document_text(document) == "# 标题\n\n内容"


def test_parse_document_text_rejects_unsupported_extension(tmp_path: Path) -> None:
    document = tmp_path / "sample.csv"
    document.write_text("value", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_document_text(document)


def test_assign_parent_orders_builds_nested_tree() -> None:
    sections = assign_parent_orders(
        [
            ("第一章", 1, "章导读"),
            ("1.1 背景", 2, "背景正文"),
            ("1.1.1 定义", 3, "定义正文"),
            ("1.2 方法", 2, "方法正文"),
            ("第二章", 1, ""),
        ]
    )

    assert [section.order for section in sections] == [0, 1, 2, 3, 4]
    assert [section.parent_order for section in sections] == [None, 0, 1, 0, None]
    assert sections[2].content == "定义正文"


def test_assign_parent_orders_tolerates_skipped_levels() -> None:
    sections = assign_parent_orders([("第一章", 1, ""), ("1.1.1 细节", 3, "x")])

    assert sections[1].parent_order == 0


def test_assign_parent_orders_empty_input() -> None:
    assert assign_parent_orders([]) == []


def test_parsed_document_defaults_to_empty() -> None:
    parsed = ParsedDocument()

    assert parsed.sections == []
    assert parsed.tables == []
    assert parsed.lead_content == ""


def test_parsed_section_rejects_level_zero() -> None:
    with pytest.raises(ValidationError):
        ParsedSection(title="x", level=0, order=0, content="")


def test_markdown_structure_parses_nested_headings() -> None:
    text = (
        "# 第一章\n\n导读\n\n## 1.1 背景\n\n背景正文\n\n"
        "### 细节\n\n细则\n\n## 1.2 方法\n\n方法正文\n"
    )

    parsed = _parse_markdown_structure(text)

    assert [section.title for section in parsed.sections] == [
        "第一章",
        "1.1 背景",
        "细节",
        "1.2 方法",
    ]
    assert [section.level for section in parsed.sections] == [1, 2, 3, 2]
    assert [section.parent_order for section in parsed.sections] == [None, 0, 1, 0]
    assert parsed.sections[0].content == "导读"
    assert parsed.lead_content == ""


def test_markdown_structure_ignores_fenced_code() -> None:
    parsed = _parse_markdown_structure("# 真标题\n\n```python\n# 不是标题\n```\n\n正文\n")

    assert [section.title for section in parsed.sections] == ["真标题"]
    assert "不是标题" in parsed.sections[0].content


def test_markdown_structure_supports_setext_and_nospace_atx() -> None:
    parsed = _parse_markdown_structure("总览\n===\n\n#紧凑标题\n\n内容\n")

    assert [(section.title, section.level) for section in parsed.sections] == [
        ("总览", 1),
        ("紧凑标题", 1),
    ]


def test_markdown_structure_collects_lead_content() -> None:
    parsed = _parse_markdown_structure("前言第一段。\n\n# 第一章\n\n正文\n")

    assert parsed.lead_content == "前言第一段。"
    assert parsed.sections[0].content == "正文"


def test_markdown_structure_without_headings_returns_empty_sections() -> None:
    parsed = _parse_markdown_structure("只有正文，没有标题。\n\n第二段。\n")

    assert parsed.sections == []
    assert parsed.lead_content == "只有正文，没有标题。\n\n第二段。"


def test_match_heading_style_handles_chinese_and_english_names() -> None:
    assert match_heading_style("Heading 1") == 1
    assert match_heading_style("标题 2") == 2
    assert match_heading_style("标题9") == 9
    assert match_heading_style("Title") == 1
    assert match_heading_style("Normal") is None
    assert match_heading_style("Heading 10") is None
    assert match_heading_style(None) is None
    assert match_heading_style("") is None


def test_docx_structure_parses_mixed_styles(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "lesson.docx"
    document = DocxDocument()
    document.add_paragraph("课程导读")
    document.add_paragraph("第一章", style="Heading 1")
    document.add_paragraph("章导读")
    document.add_paragraph("1.1 背景", style="Heading 2")
    document.add_paragraph("背景正文")
    document.save(path)

    parsed = _parse_docx_structure(path)

    assert [section.title for section in parsed.sections] == ["第一章", "1.1 背景"]
    assert [section.parent_order for section in parsed.sections] == [None, 0]
    assert parsed.lead_content == "课程导读"
    assert parsed.sections[0].content == "章导读"


def test_docx_structure_matches_chinese_style_name(tmp_path: Path) -> None:
    from docx import Document as DocxDocument
    from docx.enum.style import WD_STYLE_TYPE

    path = tmp_path / "chinese.docx"
    document = DocxDocument()
    style = document.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("第一章", style=style)
    document.add_paragraph("正文")
    document.save(path)

    parsed = _parse_docx_structure(path)

    assert [(section.title, section.level) for section in parsed.sections] == [("第一章", 1)]


def test_docx_structure_detects_bold_heuristic_heading(tmp_path: Path) -> None:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    path = tmp_path / "heuristic.docx"
    document = DocxDocument()
    document.add_paragraph("第一章", style="Heading 1")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("加粗小节")
    run.bold = True
    run.font.size = Pt(16)
    document.add_paragraph("小节正文")
    document.save(path)

    parsed = _parse_docx_structure(path)

    assert [(section.title, section.level) for section in parsed.sections] == [
        ("第一章", 1),
        ("加粗小节", 2),
    ]
    assert parsed.sections[1].parent_order == 0


def test_docx_structure_rejects_bold_non_headings(tmp_path: Path) -> None:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    path = tmp_path / "bold-noise.docx"
    document = DocxDocument()
    long_bold = document.add_paragraph()
    long_run = long_bold.add_run(
        "这是一段很长的加粗文字它的长度明显超过四十个字符"
        "因此不应该被识别为标题而应该被当作普通正文处理才对"
    )
    long_run.bold = True
    plain = document.add_paragraph()
    plain_run = plain.add_run("未加粗的大字号文字")
    plain_run.font.size = Pt(18)
    table = document.add_table(rows=1, cols=1)
    caption = table.cell(0, 0).paragraphs[0]
    caption_run = caption.add_run("表格内的加粗说明")
    caption_run.bold = True
    caption_run.font.size = Pt(16)
    document.save(path)

    parsed = _parse_docx_structure(path)

    assert parsed.sections == []
    assert "表格内的加粗说明" in parsed.lead_content


def test_pptx_structure_creates_one_section_per_slide(tmp_path: Path) -> None:
    from pptx import Presentation

    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    title_layout = presentation.slide_layouts[1]
    first = presentation.slides.add_slide(title_layout)
    first.shapes.title.text = "课程目标"
    first.placeholders[1].text = "掌握 RAG 原理"
    second = presentation.slides.add_slide(title_layout)
    second.placeholders[1].text = "只有正文没有标题"
    presentation.save(path)

    parsed = parse_document_structure(path)

    assert [(section.title, section.level) for section in parsed.sections] == [
        ("课程目标", 1),
        ("Slide 2", 1),
    ]
    assert parsed.sections[0].content == "课程目标\n掌握 RAG 原理"
    assert parsed.lead_content == ""


def test_parse_structure_dispatches_markdown(tmp_path: Path) -> None:
    document = tmp_path / "note.md"
    document.write_text("# 第一章\n\n正文\n", encoding="utf-8")

    parsed = parse_document_structure(document)

    assert [section.title for section in parsed.sections] == ["第一章"]


def test_parse_structure_falls_back_to_single_root(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    text_file = tmp_path / "plain.txt"
    text_file.write_text("第一段\n\n第二段", encoding="utf-8")
    pdf_file = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf_file.open("wb") as handle:
        writer.write(handle)

    text_parsed = parse_document_structure(text_file)
    pdf_parsed = parse_document_structure(pdf_file)

    assert text_parsed.sections == []
    assert text_parsed.lead_content == "第一段\n\n第二段"
    assert pdf_parsed.sections == []
    assert pdf_parsed.lead_content == ""


def test_parse_structure_rejects_unsupported_extension(tmp_path: Path) -> None:
    document = tmp_path / "sample.csv"
    document.write_text("value", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_document_structure(document)
